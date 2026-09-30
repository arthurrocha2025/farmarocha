"""Gera as duas planilhas finais:

1. Controle_Compras_AAAAMMDD.xlsx — Resumo, tabela de preços de cada distribuidor
   (com o que foi comprado marcado), Comprados e Pendências.
2. Pedido_AAAAMMDD.xlsx — pedido de cada distribuidor, pronto para envio.

Uso: python3 gerar_planilhas.py LISTA.xlsx COMPRE_MAIS.xlsx PROMO_REDE.xls [--desconto 5] [--data AAAAMMDD]
"""
import argparse
import datetime as dt
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

import montar_pedido_compre_mais as cm
import montar_pedido_promo_rede as pr

AQUI = Path(__file__).parent
VERDE = PatternFill("solid", fgColor="C6EFCE")
NEGRITO = Font(bold=True)


def formatar(w, destacar=None):
    for ws in w.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = NEGRITO
        cab = [c.value for c in ws[1]]
        for col in ws.columns:
            larg = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(larg + 2, 8), 45)
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        if destacar and destacar in cab:
            i = cab.index(destacar)
            for row in ws.iter_rows(min_row=2):
                if row[i].value:
                    for c in row:
                        c.fill = VERDE
        for row in ws.iter_rows(min_row=2):
            if any(c.value == "TOTAL" for c in row):
                for c in row:
                    c.font = NEGRITO


def com_total(df, qtd, total):
    t = {c: None for c in df.columns}
    t[df.columns[2]] = "TOTAL"
    t[qtd] = df[qtd].sum()
    t[total] = round(df[total].sum(), 2)
    return pd.concat([df.astype(object), pd.DataFrame([t])], ignore_index=True)


def main(lista, cm_path, promo_path, desconto, data):
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        cm.main(lista, cm_path, tmp / f"Pedido_COMPRE_MAIS_{data}.xlsx")
        pr.main(lista, promo_path, tmp / f"Pedido_PROMO_REDE_{data}.xlsx", desconto)
        a_cm = pd.read_excel(tmp / f"Pedido_COMPRE_MAIS_{data}.xlsx", sheet_name=None)
        a_pr = pd.read_excel(tmp / f"Pedido_PROMO_REDE_{data}.xlsx", sheet_name=None)

    # ---------- Tabelas de preços com o que foi comprado ----------
    tabelas = {}
    ped_cm = a_cm["Pedido"]
    comp_cm = a_cm["Comparativo"]
    for aba, fil in cm.FILIAIS.items():
        t = pd.read_excel(cm_path, sheet_name=aba)
        t["ean"] = t["CODAUXILIAR"].map(cm.norm_ean)
        c = comp_cm[comp_cm["Filial CM"] == fil].copy()
        c["ean"] = c["EAN"].map(cm.norm_ean)
        c = c.drop_duplicates("ean")[["ean", "Cód. interno", "Produto", "Un/Cx", "Preço CM (un)",
                                      "Últ. compra (un)", "Situação"]]
        p = ped_cm[ped_cm["Filial CM"] == fil][["Cód. CM", "Qtd pedido (cx)", "Total CM", "Economia"]]
        t = (t.merge(c, on="ean", how="left")
             .merge(p, left_on="COD_PROD", right_on="Cód. CM", how="left")
             .drop(columns=["ean", "Cód. CM"]))
        t = t.rename(columns={"Produto": "Produto (nossa lista)", "Preço CM (un)": "Preço (un)",
                              "Últ. compra (un)": "Nossa últ. compra (un)",
                              "Qtd pedido (cx)": "COMPRADO (cx)", "Total CM": "Total comprado"})
        tabelas[f"COMPRE MAIS {fil}"] = t

    t = pd.read_excel(promo_path, dtype={"EAN": str})
    t["ean"] = t["EAN"].map(cm.norm_ean)
    c = a_pr["Comparativo"].copy()
    c["ean"] = c["ean"].map(cm.norm_ean)
    c = c.drop_duplicates("ean")[["ean", "Cód. interno", "Produto", "Un/Cx", "Preço c/ desc (un)",
                                  "Últ. preço compra", "Situação"]]
    p = a_pr["Pedido"][["Cód. PROMO", "Qtd pedido (cx)", "Total PROMO", "Economia"]]
    t.insert(t.columns.get_loc("PRECO_FINAL") + 1, f"PRECO -{desconto:g}%",
             (t["PRECO_FINAL"] * (1 - desconto / 100)).round(2))
    t = (t.merge(c, on="ean", how="left")
         .merge(p, left_on="CODPROD", right_on="Cód. PROMO", how="left")
         .drop(columns=["ean", "Cód. PROMO"]))
    t = t.rename(columns={"Produto": "Produto (nossa lista)", "Preço c/ desc (un)": f"Preço -{desconto:g}% (un)",
                          "Últ. preço compra": "Nossa últ. compra (un)",
                          "Qtd pedido (cx)": "COMPRADO (cx)", "Total PROMO": "Total comprado"})
    tabelas["PROMO REDE"] = t

    # ---------- Comprados ----------
    comprados = pd.concat([
        ped_cm.assign(Distribuidor="COMPRE MAIS " + ped_cm["Filial CM"]).rename(columns={
            "Cód. CM": "Cód. distribuidor", "EAN CM": "EAN", "Descrição CM": "Descrição distribuidor",
            "Preço CM (cx)": "Preço (cx)", "Preço CM (un)": "Preço (un)", "Total CM": "Total"}),
        a_pr["Pedido"].assign(Distribuidor="PROMO REDE").rename(columns={
            "Cód. PROMO": "Cód. distribuidor", "EAN PROMO": "EAN", "Descrição PROMO": "Descrição distribuidor",
            f"Preço -{desconto:g}% (cx)": "Preço (cx)", f"Preço -{desconto:g}% (un)": "Preço (un)",
            "Total PROMO": "Total"}),
    ], ignore_index=True)
    comprados = comprados[["Distribuidor", "Cód. distribuidor", "EAN", "Descrição distribuidor", "Cód. interno",
                           "Produto (lista)", "Un/Cx", "Qtd pedido (cx)", "Preço (cx)", "Preço (un)",
                           "Últ. compra (un)", "Total", "Total pela últ. compra", "Economia",
                           "Fornecedor anterior", "Obs"]]

    # ---------- Resumo ----------
    res = (comprados.groupby("Distribuidor", sort=False)
           .agg(Itens=("Cód. interno", "count"), Caixas=("Qtd pedido (cx)", "sum"),
                **{"Valor do pedido": ("Total", "sum"), "Pela última compra": ("Total pela últ. compra", "sum"),
                   "Economia": ("Economia", "sum")}).reset_index())
    res.loc[len(res)] = ["TOTAL", res["Itens"].sum(), res["Caixas"].sum(), res["Valor do pedido"].sum(),
                         res["Pela última compra"].sum(), res["Economia"].sum()]
    res = res.round(2)
    res["Obs"] = ""
    res.loc[res["Distribuidor"] == "PROMO REDE", "Obs"] = f"Com {desconto:g}% de desconto adicional"

    # ---------- Pendências ----------
    pend = []
    for _, r in comprados[comprados["Obs"].notna() & (comprados["Obs"] != "")].iterrows():
        pend.append({"Distribuidor": r["Distribuidor"], "Cód. interno": r["Cód. interno"],
                     "Produto": r["Produto (lista)"], "Pendência": r["Obs"]})
    if "Sem estoque CM" in a_cm:
        for _, r in a_cm["Sem estoque CM"].iterrows():
            pend.append({"Distribuidor": "COMPRE MAIS " + r["Filial CM"], "Cód. interno": r["Cód. interno"],
                         "Produto": r["Produto (lista)"], "Pendência": "Mais barato, mas sem estoque"})
    if "Conferir" in a_pr:
        for _, r in a_pr["Conferir"].iterrows():
            pend.append({"Distribuidor": "PROMO REDE", "Cód. interno": r["Cód. interno"], "Produto": r["Produto"],
                         "Pendência": f"Conferir fração: Un/Cx {r['Un/Cx']}, últ. compra R$ {r['Últ. preço compra']:.2f}/un "
                                      f"x PROMO R$ {r['Preço c/ desc (un)']:.2f}/un — não pedido"})
    pend = pd.DataFrame(pend)

    controle = AQUI / f"Controle_Compras_{data}.xlsx"
    with pd.ExcelWriter(controle, engine="openpyxl") as w:
        res.to_excel(w, sheet_name="Resumo", index=False)
        for nome, t in tabelas.items():
            t.to_excel(w, sheet_name=nome[:31], index=False)
        comprados.to_excel(w, sheet_name="Comprados", index=False)
        pend.to_excel(w, sheet_name="Pendências", index=False)
        formatar(w, destacar="COMPRADO (cx)")

    # ---------- Pedido ----------
    pedido = AQUI / f"Pedido_{data}.xlsx"
    with pd.ExcelWriter(pedido, engine="openpyxl") as w:
        for fil in cm.FILIAIS.values():
            f = ped_cm[ped_cm["Filial CM"] == fil].sort_values("Descrição CM")
            env = pd.DataFrame({"COD_PROD": f["Cód. CM"], "CODAUXILIAR": f["EAN CM"].astype(str),
                                "DESCRICAO": f["Descrição CM"], "QTD": f["Qtd pedido (cx)"],
                                "PRECO": f["Preço CM (cx)"], "TOTAL": f["Total CM"].round(2)})
            com_total(env, "QTD", "TOTAL").to_excel(w, sheet_name=f"COMPRE MAIS {fil}", index=False)
        f = a_pr["Pedido"].sort_values("Descrição PROMO")
        env = pd.DataFrame({"CODPROD": f["Cód. PROMO"], "EAN": f["EAN PROMO"].astype(str),
                            "DESCRICAO": f["Descrição PROMO"], "QTD": f["Qtd pedido (cx)"],
                            "PRECO_FINAL": f["Preço tabela (cx)"],
                            f"PRECO -{desconto:g}%": f[f"Preço -{desconto:g}% (cx)"],
                            "TOTAL": f["Total PROMO"].round(2)})
        com_total(env, "QTD", "TOTAL").to_excel(w, sheet_name="PROMO REDE", index=False)
        formatar(w)

    print(f"\nGerados: {controle.name} e {pedido.name}")
    print(res.to_string(index=False))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("lista")
    ap.add_argument("compre_mais")
    ap.add_argument("promo")
    ap.add_argument("--desconto", type=float, default=5.0)
    ap.add_argument("--data", default=dt.date.today().strftime("%Y%m%d"))
    a = ap.parse_args()
    main(a.lista, a.compre_mais, a.promo, a.desconto, a.data)
