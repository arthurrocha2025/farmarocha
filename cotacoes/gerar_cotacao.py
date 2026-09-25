"""Gera a planilha de cotação para distribuidores a partir da Lista de Compra.

Cada EAN (principal ou adicional) vira uma linha própria: o distribuidor
cota pelo EAN que tiver cadastrado, sem precisar olhar colunas extras.

Uso: python gerar_cotacao.py <Lista_de_Compra.xlsx> [saida.xlsx]
"""
import sys
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

COLS_EAN = ["EAN princ.", "EAN adic. 1", "EAN adic. 2", "EAN adic. 3"]


def normaliza(ean):
    ean = "" if pd.isna(ean) else str(ean).strip()
    return ean[:-2] if ean.endswith(".0") else ean


def eans_do_produto(row):
    """EANs distintos do produto (ignora zeros à esquerda; prefere 13 dígitos)."""
    vistos = {}
    for c in COLS_EAN:
        e = normaliza(row.get(c))
        if not e:
            continue
        chave = e.lstrip("0")
        if chave not in vistos or (len(e) == 13 and len(vistos[chave]) != 13):
            vistos[chave] = e
    return list(vistos.values()) or [""]


def carrega(caminho):
    df = pd.read_excel(caminho, header=5, dtype={c: str for c in COLS_EAN})
    df = df[pd.to_numeric(df["Cód. interno"], errors="coerce").notna()]
    linhas = []
    for _, r in df.iterrows():
        for e in eans_do_produto(r):
            linhas.append({
                "cod": int(r["Cód. interno"]),
                "ean": e,
                "produto": str(r["Produto"]).strip(),
                "fabricante": "" if pd.isna(r["Fabricante"]) else str(r["Fabricante"]).strip(),
                "qtde": int(r["Caixas"]),
            })
    linhas.sort(key=lambda x: (x["produto"], x["cod"]))
    return linhas, len(df)


def gera(linhas, n_produtos, saida):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotação"

    F = "Arial"
    azul = "1565C0"
    fill_input = PatternFill("solid", fgColor="FFF59D")
    fill_head = PatternFill("solid", fgColor=azul)
    fill_zebra = PatternFill("solid", fgColor="F2F6FB")
    fina = Side(style="thin", color="BFBFBF")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    fonte = Font(name=F, size=10)

    # ---- Cabeçalho
    ws["A1"] = "DROGARIAS ROCHA — COTAÇÃO DE PREÇOS"
    ws["A1"].font = Font(name=F, size=16, bold=True, color=azul)
    ws["A2"] = f"Solicitação emitida em {date.today():%d/%m/%Y}  •  {len(linhas)} itens (EANs)"
    ws["A2"].font = Font(name=F, size=10, italic=True, color="595959")

    campos = [("Distribuidora:", "Prazo de pagamento:"),
              ("Vendedor(a):", "Pedido mínimo (R$):"),
              ("Telefone / WhatsApp:", "Frete / Prazo de entrega:"),
              ("E-mail:", "Validade da proposta:")]
    for i, (esq, dir_) in enumerate(campos, start=4):
        for col_lbl, col_val, txt in (("A", "B", esq), ("F", "G", dir_)):
            ws[f"{col_lbl}{i}"] = txt
            ws[f"{col_lbl}{i}"].font = Font(name=F, size=10, bold=True)
            ws[f"{col_lbl}{i}"].alignment = Alignment(horizontal="right")
        ws.merge_cells(f"B{i}:D{i}")
        ws.merge_cells(f"G{i}:I{i}")
        for c in ("B", "C", "D", "G", "H", "I"):
            ws[f"{c}{i}"].fill = fill_input
            ws[f"{c}{i}"].border = borda
            ws[f"{c}{i}"].font = fonte

    instr = [
        "COMO PREENCHER:  preencha somente as células AMARELAS (dados acima e colunas G, H e I da tabela).",
        "• Cada linha é um EAN. Informe o PREÇO LÍQUIDO (já com desconto e impostos/ST) da embalagem que o EAN representa (caixa, display etc.). Ex.: 12,34",
        "• Sem o item ou sem estoque: deixe o preço em branco. Estoque parcial: informe a quantidade disponível na coluna H.",
        "• Não altere, exclua nem reordene linhas e colunas. Devolva este mesmo arquivo em Excel (.xlsx).",
    ]
    for i, t in enumerate(instr, start=9):
        ws[f"A{i}"] = t
        ws[f"A{i}"].font = Font(name=F, size=10, bold=(i == 9), color="C62828" if i == 9 else "000000")

    # ---- Tabela
    H = 14
    cab = ["Item", "Cód. Rocha", "EAN", "Produto", "Fabricante", "Qtde (emb.)",
           "Preço líquido por emb. (R$)", "Qtde disponível (emb.)", "Observação",
           "Total (R$)"]
    larg = [7, 11, 16, 52, 22, 10, 16, 13, 28, 14]
    for j, (t, w) in enumerate(zip(cab, larg), start=1):
        c = ws.cell(row=H, column=j, value=t)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = fill_head if j not in (7, 8, 9) else PatternFill("solid", fgColor="F9A825")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[H].height = 32
    ws.cell(row=H, column=2).comment = Comment(
        "Código interno da Drogarias Rocha. Mesmo código em mais de uma linha = "
        "mesmo produto com EANs diferentes; cote o(s) EAN(s) que tiver.", "Rocha")

    ini = H + 1
    cod_ant, zebra = None, False
    for k, l in enumerate(linhas):
        r = ini + k
        if l["cod"] != cod_ant:
            zebra, cod_ant = not zebra, l["cod"]
        vals = [k + 1, l["cod"], l["ean"], l["produto"], l["fabricante"], l["qtde"], None, None, None,
                f'=IF(G{r}="","",F{r}*G{r})']
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = fonte
            c.border = borda
            if j in (7, 8, 9):
                c.fill = fill_input
            elif zebra:
                c.fill = fill_zebra
        ws.cell(row=r, column=3).number_format = "@"
        ws.cell(row=r, column=7).number_format = '#,##0.00'
        ws.cell(row=r, column=8).number_format = '0'
        ws.cell(row=r, column=10).number_format = '#,##0.00'
        for j in (1, 2, 3, 6, 8):
            ws.cell(row=r, column=j).alignment = Alignment(horizontal="center")
    fim = ini + len(linhas) - 1

    # ---- Totais
    t = fim + 2
    ws[f"I{t}"] = "Itens cotados:"
    ws[f"J{t}"] = f'=COUNT(G{ini}:G{fim})'
    ws[f"I{t+1}"] = "Total cotado (R$):"
    ws[f"J{t+1}"] = f'=SUM(J{ini}:J{fim})'
    for rr in (t, t + 1):
        ws[f"I{rr}"].font = Font(name=F, size=10, bold=True)
        ws[f"I{rr}"].alignment = Alignment(horizontal="right")
        ws[f"J{rr}"].font = Font(name=F, size=10, bold=True)
        ws[f"J{rr}"].border = borda
    ws[f"J{t+1}"].number_format = '#,##0.00'

    # Validação: preço e disponível devem ser números >= 0
    dv_p = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          errorTitle="Valor inválido",
                          error="Informe apenas o preço em número (ex.: 12,34).")
    dv_q = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                          errorTitle="Valor inválido", error="Informe a quantidade em número inteiro de embalagens.")
    ws.add_data_validation(dv_p)
    ws.add_data_validation(dv_q)
    dv_p.add(f"G{ini}:G{fim}")
    dv_q.add(f"H{ini}:H{fim}")

    # Destaca disponível menor que o pedido
    ws.conditional_formatting.add(
        f"H{ini}:H{fim}",
        FormulaRule(formula=[f'AND(H{ini}<>"",H{ini}<F{ini})'], font=Font(color="C62828", bold=True)))

    ws.freeze_panes = f"A{ini}"
    ws.auto_filter.ref = f"A{H}:J{fim}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{H}:{H}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
    print(f"{n_produtos} produtos -> {len(linhas)} linhas (EANs). Salvo em {saida}")


if __name__ == "__main__":
    origem = sys.argv[1]
    saida = sys.argv[2] if len(sys.argv) > 2 else f"Cotacao_Rocha_{date.today():%Y%m%d}.xlsx"
    linhas, n = carrega(origem)
    gera(linhas, n, saida)
