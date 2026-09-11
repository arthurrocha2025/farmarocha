# -*- coding: utf-8 -*-
"""
Verificação de preços Panpharma (ePan PRO 10/09 21:02, conta Farma Rocha 712093)
contra os preços usados nos pedidos fechados (CSV precos_712093.csv de 10/09):
  - PEDIDO_EPAN_MEDICAMENTOS_E_CRITICOS_10-09-2026.xlsx (envio imediato)
  - PEDIDO_EPAN_DEMAIS_10-09-2026.xlsx (aguardando liberação)
Só verificação — nenhum pedido é alterado.
"""
import openpyxl, re, os, math

BASE = os.path.dirname(os.path.abspath(__file__))
UP = "/root/.claude/uploads/506daea3-7102-5cb6-9758-517151fad432/"

def digits(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = re.sub(r"\D", "", str(v))
    return s if len(s) >= 7 else None

def ean13_check(body12):
    s = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(body12))
    return str((10 - s % 10) % 10)

def keys(v):
    s = digits(v)
    if not s: return []
    out = [s.lstrip("0")]
    if len(s) == 14 and s[0] != "0":
        alt = (s[1:13] + ean13_check(s[1:13])).lstrip("0")
        if alt not in out: out.append(alt)
    return out

def num(v):
    if v in (None, ""): return None
    try: f = float(str(v).replace(",", "."))
    except ValueError: return None
    return f if f > 0 else None

# ---- painel novo (aba Farma Rocha = conta 712093, a mesma do CSV) ----
wb = openpyxl.load_workbook(UP + "ba89a8ad-ePanPRO_Precos_20260910_2102.xlsx", data_only=True)
ws = wb["Farma Rocha"]
by_cod, by_ean = {}, {}
n_com_preco = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    ean, cod = r[0], r[1]
    if cod in (None, ""): continue
    precos = [p for p in (num(r[7]), num(r[9])) if p]
    preco = min(precos) if precos else None
    est = num(r[6]) or 0
    rec = dict(preco=preco, vista=num(r[7]), prazo=num(r[9]), estoque=est, prod=r[2])
    by_cod.setdefault(str(cod).strip(), rec)
    for k in keys(ean): by_ean.setdefault(k, rec)
    if preco: n_com_preco += 1
print(f"painel novo: {len(by_cod)} códigos, {n_com_preco} com preço")

# ---- pedidos fechados ----
PEDIDOS = [("MEDICAMENTOS + CRÍTICOS (envio imediato)", "PEDIDO_EPAN_MEDICAMENTOS_E_CRITICOS_10-09-2026.xlsx"),
           ("DEMAIS (aguardando liberação)", "PEDIDO_EPAN_DEMAIS_10-09-2026.xlsx")]
linhas = []   # (pedido, cod_int, cod_forn, ean, prod, emb, qtd, ult, med, preco_old, preco_new, estoque)
for tag, fn in PEDIDOS:
    wbo = openpyxl.load_workbook(os.path.join(BASE, "saida3", fn), data_only=False)
    wso = wbo.active
    hdr = [c.value for c in wso[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    for r in wso.iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[ix["Produto"]] in (None, "TOTAL"): continue
        if str(r[0]).startswith(("Pedido", "Cotação")): continue
        cod_forn = str(r[ix["Cód. forn."]]).strip() if r[ix["Cód. forn."]] not in (None, "") else None
        ean = r[ix["EAN"]]
        hit = by_cod.get(cod_forn) if cod_forn else None
        via = "cód"
        if hit is None:
            hit = next((by_ean[k] for k in keys(ean) if k in by_ean), None)
            via = "EAN"
        emb = r[ix["Un/Emb"]] or 1
        linhas.append(dict(pedido=tag, cod=r[0], cod_forn=cod_forn, ean=str(ean or ""),
                           prod=r[ix["Produto"]], emb=emb, qtd=r[ix["Qtd pedido"]] or 0,
                           preco_old=float(r[ix["Preço (emb.)"]]),
                           ult=r[ix["Últ. compra"]] if isinstance(r[ix["Últ. compra"]], (int, float)) else None,
                           med=r[ix["Média hist."]] if isinstance(r[ix["Média hist."]], (int, float)) else None,
                           novo=hit, via=via if hit else None))

# ---- classificação ----
TOL = 0.005
for L in linhas:
    h, po = L["novo"], L["preco_old"]
    if h is None:
        L["cls"], L["preco_new"], L["dif"] = "SUMIU DO PAINEL", None, None
    elif h["preco"] is None:
        L["cls"], L["preco_new"], L["dif"] = "SEM PREÇO P/ 712093", None, None
    else:
        pn = h["preco"]; L["preco_new"] = pn; L["dif"] = pn / po - 1
        if abs(pn - po) <= TOL: L["cls"] = "MANTEVE"
        elif pn < po: L["cls"] = "CAIU"
        else:
            # ainda passa no teto (4% vs últ OU 6% vs média), sobre o preço unitário?
            pun = pn / (L["emb"] or 1)
            ult, med = L["ult"], L["med"]
            ok = ((ult is not None and pun <= ult * 1.04 + 1e-4) or
                  (med is not None and pun <= med * 1.06 + 1e-4) or
                  (ult is None and med is None))
            L["cls"] = "SUBIU (dentro do teto)" if ok else "SUBIU (ESTOURA O TETO)"
    if L["novo"] is not None and (L["novo"]["estoque"] or 0) < (L["qtd"] or 0):
        L["obs_est"] = f"estoque agora {int(L['novo']['estoque'] or 0)} (pedido {int(L['qtd'])})"
    else:
        L["obs_est"] = ""

# ---- relatório console ----
from collections import Counter
for tag, _ in PEDIDOS:
    sub = [L for L in linhas if L["pedido"] == tag]
    c = Counter(L["cls"] for L in sub)
    tot_old = sum(L["preco_old"] * L["qtd"] for L in sub)
    tot_new = sum((L["preco_new"] if L["preco_new"] is not None else L["preco_old"]) * L["qtd"] for L in sub)
    print(f"\n== {tag}: {len(sub)} itens | total antigo R$ {tot_old:,.2f} -> recotado R$ {tot_new:,.2f}")
    for k, v in c.most_common(): print(f"   {k}: {v}")
    for L in sorted([x for x in sub if x["cls"].startswith("SUBIU")], key=lambda x: -(x["dif"] or 0))[:15]:
        print(f"   ^ {str(L['prod'])[:42]:42s} {L['preco_old']:8.2f} -> {L['preco_new']:8.2f} ({L['dif']:+.1%}) {L['cls']}")
    for L in [x for x in sub if x["cls"] in ("SUMIU DO PAINEL", "SEM PREÇO P/ 712093")]:
        print(f"   ? {str(L['prod'])[:42]:42s} {L['preco_old']:8.2f} {L['cls']}")
    est = [x for x in sub if x["obs_est"]]
    if est:
        print(f"   estoque insuficiente agora: {len(est)}")

import json
json.dump([{k: v for k, v in L.items() if k != "novo"} | {"estoque": (L["novo"] or {}).get("estoque")} for L in linhas],
          open(os.path.join(BASE, "verif_epan.json"), "w"), ensure_ascii=False, default=str)
print("\nok")
