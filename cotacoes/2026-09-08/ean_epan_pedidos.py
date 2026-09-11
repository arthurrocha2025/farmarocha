# -*- coding: utf-8 -*-
"""
Troca a coluna EAN dos 2 pedidos ePan pelo EAN CONFIRMADO no catálogo Panpharma
(painel ePan PRO 10/09/2026 21:02, conta Farma Rocha 712093), casado pelo
Cód. forn. (código Panpharma). Fallback: CSV precos_712093. Nomes mantidos.
"""
import openpyxl, csv, re, os

BASE = os.path.dirname(os.path.abspath(__file__))
UP = "/root/.claude/uploads/506daea3-7102-5cb6-9758-517151fad432/"
OUT = os.path.join(BASE, "saida3")

def digits(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = re.sub(r"\D", "", str(v))
    return s if len(s) >= 7 else None

# EAN por código Panpharma: painel novo (confirmado) + CSV como fallback
ean_panel, ean_csv = {}, {}
wb = openpyxl.load_workbook(UP + "ba89a8ad-ePanPRO_Precos_20260910_2102.xlsx", data_only=True)
for r in wb["Farma Rocha"].iter_rows(min_row=2, values_only=True):
    if r[1] in (None, ""): continue
    e = digits(r[0])
    if e: ean_panel.setdefault(str(r[1]).strip(), e)
with open(UP + "ddb963c0-precos_712093.csv", encoding="utf-8-sig", newline="") as fh:
    for row in csv.DictReader(fh, delimiter=";"):
        cod = str(row.get("CODIGO") or "").strip()
        e = digits(row.get("EAN"))
        if cod and e: ean_csv.setdefault(cod, e)
print(f"painel: {len(ean_panel)} códigos com EAN | csv: {len(ean_csv)}")

FILES = ["PEDIDO_EPAN_MEDICAMENTOS_E_CRITICOS_10-09-2026.xlsx", "PEDIDO_EPAN_DEMAIS_10-09-2026.xlsx"]
for fn in FILES:
    p = os.path.join(OUT, fn)
    wbo = openpyxl.load_workbook(p)   # preserva fórmulas
    ws = wbo.active
    assert ws["C1"].value == "EAN"
    ws["C1"].value = "EAN ePan (confirmado)"
    n = trocado = igual = sem = 0
    for row in ws.iter_rows(min_row=2):
        a, c = row[0].value, row[2]
        if a is None or row[3].value in (None, "TOTAL") or str(a).startswith("Pedido"): continue
        n += 1
        cod_forn = str(row[1].value).strip()
        novo = ean_panel.get(cod_forn) or ean_csv.get(cod_forn)
        antigo = digits(c.value)
        if not novo:
            sem += 1
            obs = row[16]
            obs.value = ("; ".join(x for x in [str(obs.value or "").strip(), "EAN ePan não encontrado — mantido EAN interno"] if x))
            continue
        if antigo and novo.lstrip("0") == antigo.lstrip("0"): igual += 1
        else: trocado += 1
        c.value = novo
        c.number_format = "@"
    wbo.save(p)
    print(f"{fn}: {n} itens | EAN igual ao interno: {igual} | substituído: {trocado} | sem EAN ePan: {sem}")
