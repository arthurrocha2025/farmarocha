# -*- coding: utf-8 -*-
"""
Pedido DIMEC × Radar Pareto 08/09/2026.
- Casamento por TODOS os EANs (Radar adic. 1-5 + controle 04/09 princ/adic. 1-3),
  com equivalência DUN-14 <-> EAN-13.
- SEM limiar de preços: todo item casado entra; histórico é só referência.
- Conversão de embalagem: multiplicador detectado pelo nome Dimec ((20X10), C/48 FLAC,
  C/12 UND...) ancorado no preço histórico; qtd do pedido = teto(caixas / mult).
"""
import openpyxl, re, os, math
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

UP = "/root/.claude/uploads/506daea3-7102-5cb6-9758-517151fad432/"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "saida2")
os.makedirs(OUT, exist_ok=True)

def digits(v):
    if v is None:
        return None
    s = re.sub(r"\D", "", str(v))
    return s if len(s) >= 7 else None

def ean13_check(body12):
    s = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(body12))
    return str((10 - s % 10) % 10)

def keys(v):
    s = digits(v)
    if not s:
        return set()
    out = {s.lstrip("0")}
    if len(s) == 14 and s[0] != "0":
        out.add((s[1:13] + ean13_check(s[1:13])).lstrip("0"))
    return out

def pack_candidates(nome):
    """Multiplicadores plausíveis a partir do nome Dimec."""
    cands = {1}
    for m in re.finditer(r"\((\d+)\s*X\s*\d+\)", str(nome), re.I):
        cands.add(int(m.group(1)))
    for m in re.finditer(r"C/\s*(\d+)\s*(UND|FLAC|BL)", str(nome), re.I):
        cands.add(int(m.group(1)))
    return cands

# ---- Dimec ----
wbd = openpyxl.load_workbook(UP + "f01e600f-pre_o_dimec.xlsx", data_only=True)
dimec_by_key = {}
for r in wbd["geral"].iter_rows(min_row=2, values_only=True):
    if not isinstance(r[4], (int, float)) or r[4] <= 0:
        continue
    rec = dict(ean=str(r[0]), cod=r[1], prod=r[2], marca=r[3], preco=float(r[4]),
               estoque=r[5] or 0, emb=r[6])
    for k in keys(r[0]):
        dimec_by_key.setdefault(k, rec)

# ---- Radar + EANs do controle 04/09 ----
wb = openpyxl.load_workbook(UP + "f22e9ec6-Radar_Pareto_20260908_102517.xlsx", data_only=True)
ws = wb["Radar_Pareto"]
hdr = [c.value for c in ws[6]]
wbc = openpyxl.load_workbook(UP + "b26d7f19-LISTA_DE_COMPRA_04092026__CONTROLE.xlsx", data_only=True)
wsc = wbc["Lista_de_Compra"]
hc = [c.value for c in wsc[1]]
old = {}
for r in wsc.iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    d = dict(zip(hc, r))
    old[d["Cód. interno"]] = [d.get("EAN princ.")] + [d.get(f"EAN adic. {i}") for i in range(1, 4)]

itens = []
for r in ws.iter_rows(min_row=7, values_only=True):
    if r[1] is None:
        continue
    d = dict(zip(hdr, r))
    if d["OL"] not in (None, ""):
        continue
    ks = set()
    for src in [d.get(f"EAN adic. {i}") for i in range(1, 6)] + old.get(d["Cód"], []):
        ks |= keys(src)
    hit = next((dimec_by_key[k] for k in ks if k in dimec_by_key), None)
    if not hit:
        continue
    ult = d.get("Últ. preço compra") if isinstance(d.get("Últ. preço compra"), (int, float)) and d["Últ. preço compra"] > 0 else None
    refs = [d[k] for k in ("Melhor 3m", "Melhor 6m", "Melhor 12m") if isinstance(d.get(k), (int, float)) and d[k] > 0]
    hist = min(refs) if refs else None
    ref = ult if ult is not None else hist
    cx = d.get("Caixas") or 0
    # multiplicador de embalagem ancorado na referência
    mult = 1
    if ref:
        mult = min(pack_candidates(hit["prod"]), key=lambda c: abs(hit["preco"] / c - ref))
    qtd = math.ceil(cx / mult) if mult > 1 else cx
    obs = []
    if mult > 1:
        obs.append(f"1 emb. Dimec = {mult} un do Radar (pedido = {qtd} emb.)")
    if hit["estoque"] < qtd:
        obs.append(f"estoque parcial ({int(hit['estoque'])})")
    itens.append((d, hit, ult, hist, cx, mult, qtd, "; ".join(obs)))

# ---- planilha ----
wbo = openpyxl.Workbook()
wso = wbo.active
wso.title = "Pedido"
HEAD = ["Cód. Dimec", "Cód. interno", "EAN Dimec", "Produto (Dimec)", "Produto (Radar)", "Marca",
        "Un/Emb", "Caixas Radar", "Qtd pedido (emb.)", "Preço Dimec (emb.)", "Total",
        "Preço un. equiv.", "Últ. compra", "Menor hist.", "Dif. %", "Estoque Dimec", "OBS"]
wso.append(HEAD)
itens.sort(key=lambda t: str(t[0]["Produto"]))
for i, (d, h, ult, hist, cx, mult, qtd, obs) in enumerate(itens, start=2):
    dif = ("=L%d/M%d-1" % (i, i)) if ult else (("=L%d/N%d-1" % (i, i)) if hist else None)
    wso.append([h["cod"], d["Cód"], h["ean"], h["prod"], d["Produto"], h["marca"],
                mult, cx, qtd, h["preco"], "=I%d*J%d" % (i, i),
                "=J%d/G%d" % (i, i), ult, hist, dif, h["estoque"], obs])
n = len(itens) + 1
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
widths = [10, 11, 15, 42, 42, 13, 8, 9, 10, 12, 11, 11, 11, 11, 9, 12, 34]
for i, w in enumerate(widths, 1):
    wso.column_dimensions[get_column_letter(i)].width = w
for c in wso[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    c.fill = H_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = THIN
wso.row_dimensions[1].height = 28
for row in wso.iter_rows(min_row=2, max_row=n, max_col=len(HEAD)):
    for c in row:
        c.font = Font(name=F, size=10)
        c.border = THIN
        if c.column in (10, 11, 12, 13, 14):
            c.number_format = "#,##0.00"
        if c.column in (7, 8, 9, 16):
            c.number_format = "0"
    row[2].number_format = "@"
    row[14].number_format = "0.0%"
wso.freeze_panes = "A2"
wso.cell(row=n + 1, column=5, value="TOTAL").font = Font(name=F, bold=True, size=10)
tc = wso.cell(row=n + 1, column=11, value="=SUM(K2:K%d)" % n)
tc.font = Font(name=F, bold=True, size=10)
tc.number_format = "#,##0.00"
lg = wso.cell(row=n + 3, column=1,
              value=("Pedido Dimec (catálogo 'preço 2%') × Radar Pareto 08/09/2026 — casado por TODOS os EANs "
                     "(Radar adic. 1-5 + controle 04/09), com equivalência DUN-14/EAN-13. SEM LIMIAR DE PREÇOS: "
                     "todos os itens casados entram; 'Últ. compra' e 'Menor hist.' (3/6/12m) são referência. "
                     "'Un/Emb' = unidades do Radar por embalagem Dimec (detectado pelo nome e preço); "
                     "'Qtd pedido' já convertida para embalagens Dimec (arredondada para cima)."))
lg.font = Font(name=F, italic=True, size=9)
out = os.path.join(OUT, "PEDIDO_DIMEC_08-09-2026.xlsx")
wbo.save(out)

tot = sum(h["preco"] * q for _, h, _, _, _, _, q, _ in itens)
print(f"itens: {len(itens)} | total do pedido: R$ {tot:,.2f}")
for d, h, ult, hist, cx, mult, qtd, obs in itens:
    pu = h["preco"] / mult
    print(f"  {str(d['Produto'])[:40]:40s} cx={cx:3g} mult={mult:3d} qtd={qtd:3d} "
          f"emb=R${h['preco']:8.2f} un=R${pu:7.2f} ult={ult or 0:7.2f} {obs}")
