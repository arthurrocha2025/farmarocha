# -*- coding: utf-8 -*-
"""
Correção do pedido Panpharma (11/09): 'altamente crítico' = classe A do Radar.
- MEDICAMENTOS_E_CRITICOS: medicamentos (RX/GENÉRICO/OTC) + críticos zerados
  CLASSE A abaixo da últ. compra (11 itens).
- Classe B e C que estavam marcados como crítico voltam para o DEMAIS.
Nomes de arquivo mantidos (substituição in-place).
"""
import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
UP = "/root/.claude/uploads/506daea3-7102-5cb6-9758-517151fad432/"
OUT = os.path.join(BASE, "saida3")

# ---- ABC do Radar ----
wbr = openpyxl.load_workbook(UP + "f22e9ec6-Radar_Pareto_20260908_102517.xlsx", data_only=True)
wsr = wbr["Radar_Pareto"]
hdr = [c.value for c in wsr[6]]
abc = {}
for r in wsr.iter_rows(min_row=7, values_only=True):
    if r[1] is None: continue
    d = dict(zip(hdr, r))
    abc[d["Cód"]] = d.get("ABC")

# ---- lê os dois pedidos atuais ----
HEAD = ["Cód. interno","Cód. forn.","EAN","Produto","Un/Emb","Caixas","Qtd pedido","Preço (emb.)","Total",
        "Preço un. eq.","Últ. compra","Menor hist.","Média hist.","% vs média hist.","% vs últ. compra",
        "Status","OBS","Categoria"]
rows = []
for fn in ("PEDIDO_EPAN_MEDICAMENTOS_E_CRITICOS_10-09-2026.xlsx", "PEDIDO_EPAN_DEMAIS_10-09-2026.xlsx"):
    wbo = openpyxl.load_workbook(os.path.join(OUT, fn))
    wso = wbo.active
    hp = [c.value for c in wso[1]]
    assert hp == HEAD, (fn, hp)
    for r in wso.iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[3] in (None, "TOTAL") or str(r[0]).startswith("Pedido"): continue
        rows.append(list(r))
assert len(rows) == 338, len(rows)

# ---- reclassifica ----
MED = {"RX", "GENERICO", "OTC"}
med, crit_a, demais, volta = [], [], [], []
for r in rows:
    e_crit = "CRÍTICO" in str(r[16] or "")
    if r[17] in MED and not e_crit:
        med.append(r)
    elif e_crit:
        classe = abc.get(r[0])
        if classe == "A":
            r[16] = "CRÍTICO ZERADO classe A — abaixo da últ. compra"
            crit_a.append(r)
        else:
            r[16] = ""  # remove a marca; item comum
            volta.append(r); demais.append(r)
    else:
        demais.append(r)

# ---- grava ----
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E79")
CRIT_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

def grava(fn, sel, legenda):
    sel = sorted(sel, key=lambda x: str(x[3]))
    wbo = openpyxl.Workbook(); ws = wbo.active; ws.title = "Pedido"
    ws.append(HEAD)
    for i, r in enumerate(sel, start=2):
        out = list(r)
        out[2] = str(out[2] or "")
        out[8] = f"=G{i}*H{i}"
        out[9] = f"=H{i}/E{i}"
        out[13] = f"=J{i}/M{i}-1" if isinstance(r[12], (int, float)) else None
        out[14] = f"=J{i}/K{i}-1" if isinstance(r[10], (int, float)) else None
        ws.append(out)
    n = len(sel) + 1
    widths = [11,10,15,46,8,8,9,11,11,11,10,10,10,10,10,17,30,10]
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = THIN
    ws.row_dimensions[1].height = 26
    for row in ws.iter_rows(min_row=2, max_row=n, max_col=len(HEAD)):
        e_crit = "CRÍTICO" in str(row[16].value or "")
        for c in row:
            c.font = Font(name=F, size=10); c.border = THIN
            if c.column in (8,9,10,11,12,13): c.number_format = "#,##0.00"
            if c.column in (5,6,7): c.number_format = "0"
            if c.column in (14,15): c.number_format = "0.0%"
            if e_crit: c.fill = CRIT_FILL
        row[2].number_format = "@"
    ws.freeze_panes = "A2"
    ws.cell(row=n+1, column=4, value="TOTAL").font = Font(name=F, bold=True, size=10)
    tc = ws.cell(row=n+1, column=9, value=f"=SUM(I2:I{n})")
    tc.font = Font(name=F, bold=True, size=10); tc.number_format = "#,##0.00"
    lg = ws.cell(row=n+3, column=1, value=legenda)
    lg.font = Font(name=F, italic=True, size=9)
    wbo.save(os.path.join(OUT, fn))
    tot = sum(float(r[7]) * float(r[6] or 0) for r in sel)
    print(f"{fn}: {len(sel)} itens, R$ {tot:,.2f}")
    return tot

t1 = grava("PEDIDO_EPAN_MEDICAMENTOS_E_CRITICOS_10-09-2026.xlsx", med + crit_a,
           ("Pedido Panpharma (ePan) — ENVIO IMEDIATO. Medicamentos (RX/GENÉRICO/OTC) + itens ALTAMENTE críticos: "
            "classe A do Radar, zerados e abaixo da última compra (linhas salmão). Cotação 08/09 fechada 09/09; "
            "teto de reajuste 4% vs últ. compra ou 6% vs média histórica. Corrigido 11/09: críticos classe B/C "
            "movidos para o pedido DEMAIS."))
t2 = grava("PEDIDO_EPAN_DEMAIS_10-09-2026.xlsx", demais,
           ("Pedido Panpharma (ePan) — DEMAIS ITENS (aguardando liberação). Cotação 08/09 fechada 09/09; "
            "teto de reajuste 4% vs últ. compra ou 6% vs média histórica. Corrigido 11/09: inclui os itens "
            "classe B/C antes marcados como crítico."))
print(f"\nmed={len(med)} + críticos A={len(crit_a)} | demais={len(demais)} (voltaram {len(volta)}) | soma {t1+t2:,.2f}")
assert len(med) + len(crit_a) + len(demais) == 338
