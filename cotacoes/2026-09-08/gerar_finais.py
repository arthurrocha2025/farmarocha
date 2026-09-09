# -*- coding: utf-8 -*-
"""
FECHAMENTO da cotação 08/09 (finalizada em 09/09):
- 1 arquivo de pedido por distribuidor (somente itens aprovados: até 5% de reajuste; 10% se zerado);
- 1 lista única com TODOS os itens de OL do Radar (ignora cotação — vai direto para os OLs).
"""
import os
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
# reusa todo o pipeline da parcial (cargas, casamento, regra de reajuste) sem regravar a parcial
src = open(os.path.join(BASE, "gerar_parcial.py")).read().split("# ---------- estilos ----------")[0]
exec(src)

import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
OUT = os.path.join(BASE, "saida3")
os.makedirs(OUT, exist_ok=True)

def style(ws, widths, nrows, money_cols=(), int_cols=(), pct_cols=()):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = THIN
    ws.row_dimensions[1].height = 26
    for row in ws.iter_rows(min_row=2, max_row=nrows, max_col=len(widths)):
        for c in row:
            c.font = Font(name=F, size=10); c.border = THIN
            if c.column in money_cols: c.number_format = "#,##0.00"
            if c.column in int_cols: c.number_format = "0"
            if c.column in pct_cols: c.number_format = "0.0%"
    ws.freeze_panes = "A2"

# ---------------- pedidos por distribuidor ----------------
PH = ["Cód. interno","Cód. forn.","EAN","Produto","Un/Emb","Caixas","Qtd pedido","Preço (emb.)","Total",
      "Preço un. eq.","Últ. compra","Menor hist.","Média hist.","% vs média hist.","% vs últ. compra",
      "Status","OBS"]
resumo = []
for f in FORNS:
    sel = [it for it in itens if it["win"] == f and it["aceito"]]
    if not sel:
        continue
    sel.sort(key=lambda x: str(x["d"]["Produto"]))
    wbo = openpyxl.Workbook(); wsP = wbo.active; wsP.title = "Pedido"
    wsP.append(PH)
    for i, it in enumerate(sel, start=2):
        w = it["q"][f]
        dmed = f"=J{i}/M{i}-1" if it["med"] else None
        dult = f"=J{i}/K{i}-1" if it["ult"] else None
        wsP.append([it["d"]["Cód"], w["cod"], str(it["ean"]), it["d"]["Produto"], w["mult"], it["cx"],
                    w["qtd"], w["preco"], f"=G{i}*H{i}", f"=H{i}/E{i}", it["ult"], it["hist"], it["med"],
                    dmed, dult, it["status"], w["obs"]])
    n = len(sel) + 1
    style(wsP, [11,10,15,46,8,8,9,11,11,11,10,10,10,10,10,17,26], n,
          money_cols=(8,9,10,11,12,13), int_cols=(5,6,7), pct_cols=(14,15))
    for row in wsP.iter_rows(min_row=2, max_row=n, min_col=3, max_col=3): row[0].number_format = "@"
    wsP.cell(row=n+1, column=4, value="TOTAL").font = Font(name=F, bold=True, size=10)
    tc = wsP.cell(row=n+1, column=9, value=f"=SUM(I2:I{n})")
    tc.font = Font(name=F, bold=True, size=10); tc.number_format = "#,##0.00"
    lg = wsP.cell(row=n+3, column=1,
                  value=(f"Pedido FINAL {f} — cotação 08/09/2026 (fechada 09/09). Somente itens aprovados: "
                         "preço até 5% acima da referência (últ. compra; sem ela, menor histórico), "
                         "10% quando o item está zerado. 'Qtd pedido' na embalagem do fornecedor."))
    lg.font = Font(name=F, italic=True, size=9)
    fn = f"PEDIDO_{f}_09-09-2026.xlsx"
    wbo.save(os.path.join(OUT, fn))
    tot = sum(it["q"][f]["preco"] * it["q"][f]["qtd"] for it in sel)
    resumo.append((f, len(sel), tot))
    print(f"{fn}: {len(sel)} itens, R$ {tot:,.2f}")

# ---------------- lista única de OL ----------------
ol_rows = []
for r in wsr.iter_rows(min_row=7, values_only=True):
    if r[1] is None: continue
    d = dict(zip(hr, r))
    if d["OL"] in (None, ""): continue
    ean = (nec_eans.get(d["Cód"]) or [""])[0]
    dt = d.get("Data últ. compra")
    ol_rows.append([str(d["OL"]).strip(), d["Cód"], str(ean), d["Produto"], d.get("Fabricante") or "",
                    d.get("Grupo") or "", d.get("ABC") or "", d.get("Caixas"), d.get("Necessidade"),
                    d.get("Últ. preço compra") or None,
                    dt.strftime("%d/%m/%Y") if hasattr(dt, "strftime") else (dt or "")])
ol_rows.sort(key=lambda x: (x[0], str(x[3])))
wbo = openpyxl.Workbook(); wsO = wbo.active; wsO.title = "Lista_OL"
wsO.append(["OL","Cód. interno","EAN princ.","Produto","Fabricante","Grupo","ABC","Caixas","Unidades",
            "Últ. preço compra","Data últ. compra"])
for row in ol_rows: wsO.append(row)
nO = len(ol_rows) + 1
style(wsO, [26,11,15,46,20,16,6,8,9,12,12], nO, money_cols=(10,), int_cols=(8,9))
for row in wsO.iter_rows(min_row=2, max_row=nO, min_col=3, max_col=3): row[0].number_format = "@"
wsO.cell(row=nO+1, column=4, value="TOTAL").font = Font(name=F, bold=True, size=10)
for col in ("H","I"):
    c = wsO.cell(row=nO+1, column=8 if col=="H" else 9, value=f"=SUM({col}2:{col}{nO})")
    c.font = Font(name=F, bold=True, size=10); c.number_format = "0"
lg = wsO.cell(row=nO+3, column=1,
              value="Todos os itens com OL do Radar Pareto 08/09 em uma única lista (a coluna OL agrupa; "
                    "sem cotação — compra direta no OL). EAN da base Necessidade 08/09; 'Últ. preço compra' apenas referência.")
lg.font = Font(name=F, italic=True, size=9)
wbo.save(os.path.join(OUT, "LISTA_OL_09-09-2026.xlsx"))
from collections import Counter
cnt_ol = Counter(r[0] for r in ol_rows)
print(f"\nLISTA_OL: {len(ol_rows)} itens em {len(cnt_ol)} OLs")
for k, v in cnt_ol.most_common(): print(f"  {k}: {v}")
tot_ped = sum(t for _, _, t in resumo)
print(f"\nTOTAL pedidos distribuidores: {sum(n for _, n, _ in resumo)} itens, R$ {tot_ped:,.2f}")
