"""Comparativo da cotação de 05/10/2026 entre os distribuidores que responderam.

Entradas aceitas por distribuidor (NOME=arquivo):
  - tabela ePan PRO (Panpharma): abas "Farma Rocha" e "Drogaria Rocha"; preço
    líquido à vista = preço 7d − desc. %, o menor entre os dois CNPJs;
  - a própria planilha COTACAO_<NOME> respondida: preço na coluna G (0 ou vazio
    = não tem), disponível na coluna H (vazio = atende tudo). Se o distribuidor
    acrescentou as colunas "PREÇO DE VENDA"/"DESCONTO"/"ESTOQUE - UNIDADES"
    (Nazária), o preço é venda × (1 − desconto) e o estoque vem dessa coluna;
  - planilha "cotacao" do portal da SB Log (cabeçalho com "Vl. Líquido"): só
    traz os itens disponíveis; preço = Vl. Líquido.

Panpharma e Nazária cobram 3,92% de imposto por fora: entra no preço.
Só vale EAN com preço e estoque. Por produto e distribuidor vale o EAN mais
barato. Preço abaixo de 50% ou acima de 200% do custo atual (Custo un × Un/Cx)
costuma ser embalagem diferente: fica marcado para conferir e não concorre.
Teto por produto = maior preço entre última compra, custo atual e melhores
3/6/12 meses (× Un/Cx): preço acima do teto não concorre (não compramos).
Vencedor = menor preço entre os que concorrem.

Uso: python comparativo.py <Lista_de_Compra.xlsx> NOME=arquivo [NOME=arquivo ...] [-o saida.xlsx]
"""
import sys
from datetime import date

import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from gerar_cotacao import carrega

LIM_BAIXO, LIM_ALTO = 0.5, 2.0
CNPJS = ["Farma Rocha", "Drogaria Rocha"]
IMPOSTO = {"PANPHARMA": 0.0392, "NAZARIA": 0.0392}
HIST = ["Últ. preço compra", "Custo un", "Melhor 3m", "Melhor 6m", "Melhor 12m"]


def le_epan(caminho):
    out = None
    for aba in CNPJS:
        p = pd.read_excel(caminho, sheet_name=aba, dtype={"EAN": str})
        p["k"] = p["EAN"].fillna("").str.strip().str.lstrip("0")
        p[aba] = (p["Compra a vista (7d)"] * (1 - p["Desc% a vista"].fillna(0) / 100)
                  ).where(p["Compra a vista (7d)"] > 0).round(2)
        cols = ["k", aba] if out is not None else ["k", "Estoque", aba]
        p = p[p["k"] != ""].drop_duplicates("k")[cols]
        out = p if out is None else out.merge(p, on="k", how="outer")
    out["preco"] = out[CNPJS].min(axis=1)
    out["obs"] = ""
    tem = out["preco"].notna()
    out.loc[tem, "obs"] = "CNPJ " + out.loc[tem, CNPJS].idxmin(axis=1)
    out["disp"] = out["Estoque"].fillna(0)
    return out[["k", "preco", "disp", "obs"]]


def le_resposta(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    ini = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "Item")
    cab = {str(c.value).strip().upper(): c.column - 1 for c in ws[ini] if c.value}
    extra = {"PREÇO DE VENDA", "DESCONTO", "ESTOQUE - UNIDADES"} <= set(cab)
    rows = []
    for r in ws.iter_rows(min_row=ini + 1, values_only=True):
        if not isinstance(r[0], int):
            break
        if extra:
            venda = pd.to_numeric(r[cab["PREÇO DE VENDA"]], errors="coerce")
            desc = pd.to_numeric(r[cab["DESCONTO"]], errors="coerce")
            preco = round(venda * (1 - (0 if pd.isna(desc) else desc)), 2)
            disp = pd.to_numeric(r[cab["ESTOQUE - UNIDADES"]], errors="coerce")
            disp = 0 if pd.isna(disp) else disp
        else:
            preco = pd.to_numeric(r[6], errors="coerce")
            disp = pd.to_numeric(r[7], errors="coerce")
        rows.append({"k": str(r[2] or "").strip().lstrip("0"),
                     "preco": preco if preco and preco > 0 else None,  # NaN > 0 é False
                     "disp": r[5] if pd.isna(disp) else disp, "obs": r[8] or ""})
    return pd.DataFrame(rows).drop_duplicates("k")


SEM_LIMITE = 10 ** 6  # distribuidor só informou que tem, sem quantidade


def le_portal(ws, ini):
    cab = {str(c.value).strip(): c.column - 1 for c in ws[ini] if c.value}
    rows = []
    for r in ws.iter_rows(min_row=ini + 1, values_only=True):
        ean = str(r[cab["EAN"]] or "").strip().lstrip("0")
        preco = pd.to_numeric(r[cab["Vl. Líquido"]], errors="coerce")
        if ean and preco and preco > 0:
            rows.append({"k": ean, "preco": float(preco), "disp": SEM_LIMITE, "obs": ""})
    return pd.DataFrame(rows).drop_duplicates("k")


def le(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if set(CNPJS) <= set(wb.sheetnames):
        return le_epan(caminho)
    ws = wb.active
    for r in range(1, 20):
        if any(c.value == "Vl. Líquido" for c in ws[r]):
            return le_portal(ws, r)
    return le_resposta(caminho)


def monta(lista, fontes):
    d = pd.DataFrame(carrega(lista)[0])
    d["k"] = d["ean"].str.lstrip("0")
    ctl = pd.read_excel(lista, header=5)
    ctl = ctl[pd.to_numeric(ctl["Cód. interno"], errors="coerce").notna()].copy()
    ctl["cod"] = ctl["Cód. interno"].astype(int)
    ctl["custo"] = (ctl["Custo un"] * ctl["Un/Cx"]).round(2)
    ctl["teto"] = (ctl[HIST].apply(pd.to_numeric, errors="coerce").max(axis=1) * ctl["Un/Cx"]).round(2)
    prod = (d.drop_duplicates("cod")[["cod", "produto", "qtde"]]
            .merge(ctl[["cod", "Grupo", "ABC", "Fornecedor", "Un/Cx", "custo", "teto"]], on="cod"))

    for nome, arq in fontes:
        x = le(arq)
        x["preco"] = (x["preco"] * (1 + IMPOSTO.get(nome, 0))).round(2)
        t = d[d["k"] != ""].merge(x, on="k")
        t = t[t["preco"].notna() & (t["disp"] > 0)].merge(prod[["cod", "custo", "teto"]], on="cod")
        t["emb"] = (t["preco"] / t["custo"]).between(LIM_BAIXO, LIM_ALTO)
        t["teto_ok"] = t["preco"] <= t["teto"] + 0.005
        t["ok"] = t["emb"] & t["teto_ok"]
        t = t.sort_values(["cod", "ok", "emb", "preco"], ascending=[True, False, False, True]).drop_duplicates("cod")
        t = t.rename(columns={c: f"{nome}|{c}" for c in ["preco", "disp", "ean", "ok", "emb", "obs"]})
        prod = prod.merge(t[["cod"] + [c for c in t.columns if c.startswith(nome + "|")]], on="cod", how="left")
        for c in ("ok", "emb"):
            prod[f"{nome}|{c}"] = prod[f"{nome}|{c}"].fillna(False).astype(bool)

    nomes = [n for n, _ in fontes]
    prod["venc"], prod["vpreco"], prod["conferir"], prod["acima"] = "", None, "", ""
    for i, r in prod.iterrows():
        cands = [(r[f"{n}|preco"], n) for n in nomes if r[f"{n}|ok"]]
        if cands:
            p, n = min(cands)
            prod.at[i, "venc"], prod.at[i, "vpreco"] = n, p
        prod.at[i, "conferir"] = ", ".join(n for n in nomes
                                          if pd.notna(r[f"{n}|preco"]) and not r[f"{n}|emb"])
        prod.at[i, "acima"] = ", ".join(n for n in nomes if r[f"{n}|emb"] and not r[f"{n}|ok"])
    return prod.sort_values("produto"), nomes


F = "Arial"
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
M = "#,##0.00"


def cabecalho(ws, cab, larg, cor="1565C0"):
    for j, (t, w) in enumerate(zip(cab, larg), start=1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=cor)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def celula(ws, r, j, v, fmt=None, fill=None, bold=False):
    v = None if (isinstance(v, float) and pd.isna(v)) else v
    c = ws.cell(row=r, column=j, value=v)
    c.font = Font(name=F, size=10, bold=bold)
    c.border = BORDA
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def gera(prod, nomes, saida):
    wb = Workbook()
    rs = wb.active
    rs.title = "Resumo"

    # ---- Comparativo
    ws = wb.create_sheet("Comparativo")
    cab = ["Cód. Rocha", "Produto", "Grupo", "ABC", "Forn. atual", "Qtde (emb.)", "Custo atual (emb.)",
           "Teto (emb.)"]
    cab += [f"{n} (c/ imposto)" if n in IMPOSTO else n for n in nomes]
    cab += ["Vencedor", "Melhor preço", "Dif. % vs custo", "Total vencedor", "Economia", "Conferir emb./EAN",
            "Acima do teto"]
    cabecalho(ws, cab, [10, 44, 14, 5, 24, 8, 11, 11] + [12] * len(nomes) + [13, 11, 9, 12, 11, 20, 20])
    for i, (_, r) in enumerate(prod.iterrows(), start=2):
        vals = [r["cod"], r["produto"], r["Grupo"], r["ABC"], r["Fornecedor"], r["qtde"], r["custo"], r["teto"]]
        for j, v in enumerate(vals, start=1):
            celula(ws, i, j, v, M if j in (7, 8) else None)
        for j, n in enumerate(nomes, start=9):
            p = r[f"{n}|preco"]
            fill = None
            if pd.notna(p):
                fill = ("C8E6C9" if n == r["venc"] else None if r[f"{n}|ok"]
                        else "FFF59D" if not r[f"{n}|emb"] else "FFCDD2")
            celula(ws, i, j, p, M, fill, bold=(n == r["venc"]))
        j = 9 + len(nomes)
        v = r["vpreco"]
        celula(ws, i, j, r["venc"] or "Sem cotação", fill=None if r["venc"] else "FFCDD2")
        celula(ws, i, j + 1, v, M)
        celula(ws, i, j + 2, (v / r["custo"] - 1) if v and r["custo"] else None, "0.0%")
        celula(ws, i, j + 3, round(v * r["qtde"], 2) if v else None, M)
        celula(ws, i, j + 4, round((r["custo"] - v) * r["qtde"], 2) if v else None, M)
        celula(ws, i, j + 5, r["conferir"], fill="FFF59D" if r["conferir"] else None)
        celula(ws, i, j + 6, r["acima"], fill="FFCDD2" if r["acima"] else None)
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cab)).column_letter}{len(prod) + 1}"

    # ---- Pedido por distribuidor
    for n in nomes:
        v = prod[prod["venc"] == n]
        wp = wb.create_sheet(f"Pedido {n}")
        cabecalho(wp, ["Cód. Rocha", "EAN", "Produto", "Qtde (emb.)", "Disponível", "Preço", "Total",
                       "Custo atual (emb.)", "Dif. %", "Obs."], [10, 15, 44, 8, 9, 10, 11, 11, 8, 26], "2E7D32")
        for i, (_, r) in enumerate(v.iterrows(), start=2):
            p = r[f"{n}|preco"]
            disp = r[f"{n}|disp"]
            vals = [(r["cod"], None), (r[f"{n}|ean"], "@"), (r["produto"], None), (r["qtde"], None),
                    (int(disp) if disp < SEM_LIMITE else None, None), (p, M), (round(p * min(r["qtde"], disp), 2), M), (r["custo"], M),
                    (p / r["custo"] - 1, "0.0%"), (r[f"{n}|obs"], None)]
            for j, (val, fmt) in enumerate(vals, start=1):
                celula(wp, i, j, val, fmt, "FFCDD2" if j == 5 and disp < r["qtde"] else None)
        t = len(v) + 2
        celula(wp, t, 6, "Total", bold=True)
        celula(wp, t, 7, f"=SUM(G2:G{t - 1})", M, bold=True)

    # ---- Sem cotação
    sc = prod[prod["venc"] == ""]
    wz = wb.create_sheet("Sem cotação")
    cabecalho(wz, ["Cód. Rocha", "Produto", "Grupo", "Forn. atual", "Qtde (emb.)", "Custo atual (emb.)",
                   "Teto (emb.)", "Acima do teto", "Menor preço acima do teto", "Cotado só c/ emb. a conferir"],
              [10, 44, 14, 26, 8, 11, 11, 20, 12, 24], "C62828")
    for i, (_, r) in enumerate(sc.iterrows(), start=2):
        for j, (val, fmt) in enumerate([(r["cod"], None), (r["produto"], None), (r["Grupo"], None),
                                        (r["Fornecedor"], None), (r["qtde"], None), (r["custo"], M),
                                        (r["teto"], M), (r["acima"], None),
                                        (min((r[f"{n}|preco"] for n in nomes if r[f"{n}|emb"]), default=None), M),
                                        (r["conferir"], None)], start=1):
            celula(wz, i, j, val, fmt)

    # ---- Resumo
    linhas = [(f"Comparativo da cotação 05/10/2026 — {', '.join(nomes)}", None, None, None),
              (f"Gerado em {date.today():%d/%m/%Y}. {len(prod)} produtos (sem OL). Preço por embalagem de compra.",
               None, None, None),
              (None, None, None, None),
              ("Distribuidor", "Produtos cotados", "Vence em", "Total pedido (R$)")]
    for n in nomes:
        v = prod[prod["venc"] == n]
        tot = sum(r[f"{n}|preco"] * min(r["qtde"], r[f"{n}|disp"]) for _, r in v.iterrows())
        linhas.append((n, int(prod[f"{n}|preco"].notna().sum()), len(v), round(tot, 2)))
    sem = prod["venc"] == ""
    linhas.append(("Sem cotação", None, int(sem.sum()), None))
    linhas.append(("  dos quais só por estar acima do teto", None, int((sem & (prod["acima"] != "")).sum()), None))
    w = prod[prod["venc"] != ""]
    cur, tot = (w["custo"] * w["qtde"]).sum(), (w["vpreco"] * w["qtde"]).sum()
    linhas += [(None, None, None, None),
               ("Itens com vencedor pelo custo atual (R$)", None, None, round(cur, 2)),
               ("Itens com vencedor pelo melhor preço (R$)", None, None, round(tot, 2)),
               ("Diferença (R$)", None, None, round(cur - tot, 2)),
               ("Itens vencedores mais caros que o custo atual", None, int((w["vpreco"] > w["custo"]).sum()), None),
               ("Itens com embalagem/EAN a conferir", None, int((prod["conferir"] != "").sum()), None),
               (None, None, None, None),
               ("Panpharma e Nazária com +3,92% de imposto. Teto = maior entre última compra, custo e "
                "melhores 3/6/12m.", None, None, None)]
    for i, row in enumerate(linhas, start=1):
        for j, v in enumerate(row, start=1):
            if v is None:
                continue
            c = rs.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=14 if i == 1 else 10, bold=i in (1, 4) or j > 1)
            if j == 4 and isinstance(v, float):
                c.number_format = M
    for col, wdt in zip("ABCD", [46, 16, 12, 18]):
        rs.column_dimensions[col].width = wdt
    wb.save(saida)
    for row in linhas[3:]:
        if row[0]:
            print(" | ".join("" if v is None else str(v) for v in row))


if __name__ == "__main__":
    args = sys.argv[1:]
    saida = f"COMPARATIVO_{date.today():%d-%m-%Y}.xlsx"
    if "-o" in args:
        k = args.index("-o")
        saida = args[k + 1]
        del args[k:k + 2]
    lista, fontes = args[0], [tuple(a.split("=", 1)) for a in args[1:]]
    gera(*monta(lista, fontes), saida)
