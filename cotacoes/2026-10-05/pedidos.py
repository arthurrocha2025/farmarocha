"""Gera os pedidos da cotação, um arquivo por OL e um por distribuidor.

- OL: todo item com a coluna OL preenchida na Lista de Compra é comprado
  direto do laboratório. Cada OL vira um pedido independente
  (pedidos/OL/OL_<NOME>_<data>.xlsx), com quantidade, último preço e teto.
- Distribuidor: os demais itens vão para cotação; o vencedor de cada produto
  vem do comparativo (mesmas regras: imposto, teto, embalagem) e cada
  distribuidor vira um pedido (pedidos/distribuidores/PEDIDO_<NOME>_<data>.xlsx).

Uso: python pedidos.py <Lista_de_Compra.xlsx> [NOME=resposta ...] [-d dd-mm-aaaa]
Sem respostas, gera só os pedidos de OL.
"""
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from comparativo import IMPOSTO, HIST, M, SEM_LIMITE, cabecalho, celula, monta
from gerar_cotacao import COLS_EAN, normaliza

AQUI = Path(__file__).resolve().parent


def nome_arquivo(s):
    return re.sub(r"_+", "_", re.sub(r"[^A-Z0-9]+", "_", s.upper())).strip("_")


def titulo(ws, linhas, n, col_rot, col_tot):
    """Insere o título acima da tabela (n linhas de dados) e a linha de total."""
    ws.insert_rows(1, len(linhas) + 1)
    for i, t in enumerate(linhas, start=1):
        ws.cell(row=i, column=1, value=t).font = Font(name="Arial", size=14 if i == 1 else 10, bold=i == 1)
    h = len(linhas) + 2
    ws.freeze_panes = f"A{h + 1}"
    letra = ws.cell(row=h, column=col_tot).column_letter
    celula(ws, h + n + 1, col_rot, "Total", bold=True)
    celula(ws, h + n + 1, col_tot, f"=SUM({letra}{h + 1}:{letra}{h + n})", M, bold=True)


def pedidos_ol(lista, data):
    ctl = pd.read_excel(lista, header=5, dtype={c: str for c in COLS_EAN})
    ctl = ctl[pd.to_numeric(ctl["Cód. interno"], errors="coerce").notna()]
    ctl = ctl[ctl["OL"].notna() & (ctl["OL"].astype(str).str.strip() != "")].copy()
    pasta = AQUI / "pedidos" / "OL"
    pasta.mkdir(parents=True, exist_ok=True)
    resumo = []
    for ol, g in ctl.groupby(ctl["OL"].str.strip()):
        wb = Workbook()
        ws = wb.active
        ws.title = "Pedido"
        cab = ["Cód. Rocha", "EAN", "EANs adicionais", "Produto", "Fabricante", "Grupo", "Qtde (cx)", "Un/Cx",
               "Unidades", "Últ. preço (cx)", "Teto (cx)", "Total estimado"]
        cabecalho(ws, cab, [10, 15, 30, 46, 20, 16, 9, 7, 9, 11, 11, 12], "6A1B9A")
        g = g.sort_values("Produto")
        for i, (_, r) in enumerate(g.iterrows(), start=2):
            eans = [normaliza(r[c]) for c in COLS_EAN]
            un = int(r["Un/Cx"])
            ult = pd.to_numeric(r["Últ. preço compra"], errors="coerce")
            ult = round(ult * un, 2) if pd.notna(ult) and ult > 0 else None
            teto = pd.to_numeric(r[HIST], errors="coerce").max()
            teto = round(teto * un, 2) if pd.notna(teto) and teto > 0 else None
            vals = [(int(r["Cód. interno"]), None), (eans[0], "@"), (", ".join(e for e in eans[1:] if e), "@"),
                    (r["Produto"], None), (r["Fabricante"] or None, None), (r["Grupo"], None),
                    (int(r["Caixas"]), None), (un, None), (int(r["Un. compradas"]), None),
                    (ult, M), (teto, M), (round((ult or teto or 0) * int(r["Caixas"]), 2) or None, M)]
            for j, (v, f) in enumerate(vals, start=1):
                celula(ws, i, j, v, f)
        titulo(ws, [f"Pedido OL — {ol}", f"Drogarias Rocha • {data} • {len(g)} itens • "
                    "preço de referência = última compra (teto = maior preço do histórico)"], len(g), 11, 12)
        saida = pasta / f"OL_{nome_arquivo(ol)}_{data}.xlsx"
        wb.save(saida)
        resumo.append((ol, len(g), saida.name))
    return resumo


def pedidos_distribuidor(prod, nomes, data):
    pasta = AQUI / "pedidos" / "distribuidores"
    pasta.mkdir(parents=True, exist_ok=True)
    resumo = []
    for n in nomes:
        v = prod[prod["venc"] == n]
        tx = IMPOSTO.get(n, 0)
        wb = Workbook()
        ws = wb.active
        ws.title = "Pedido"
        cab = ["Cód. Rocha", "EAN", "Produto", "Qtde (emb.)", "Disponível", "Preço cotado"]
        cab += ["Preço c/ imposto"] if tx else []
        cab += ["Total", "Teto (emb.)", "Obs."]
        cabecalho(ws, cab, [10, 15, 46, 9, 9, 11] + ([11] if tx else []) + [12, 11, 24], "2E7D32")
        tot = 0
        for i, (_, r) in enumerate(v.iterrows(), start=2):
            p, disp = r[f"{n}|preco"], r[f"{n}|disp"]
            q = int(min(r["qtde"], disp))
            tot += p * q
            obs = r[f"{n}|obs"] or ""
            if q < r["qtde"]:
                obs = (obs + "; " if obs else "") + f"parcial: pedido {r['qtde']}"
            vals = [(r["cod"], None), (r[f"{n}|ean"], "@"), (r["produto"], None), (q, None),
                    (int(disp) if disp < SEM_LIMITE else None, None), (round(p / (1 + tx), 2), M)]
            vals += [(p, M)] if tx else []
            vals += [(round(p * q, 2), M), (r["teto"], M), (obs, None)]
            for j, (val, f) in enumerate(vals, start=1):
                celula(ws, i, j, val, f)
        extra = f" • total inclui {tx:.2%} de imposto".replace(".", ",") if tx else ""
        titulo(ws, [f"Pedido — {n}", f"Drogarias Rocha • cotação {data} • {len(v)} itens{extra}"],
               len(v), len(cab) - 3, len(cab) - 2)
        saida = pasta / f"PEDIDO_{nome_arquivo(n)}_{data}.xlsx"
        wb.save(saida)
        resumo.append((n, len(v), round(tot, 2), saida.name))
    return resumo


if __name__ == "__main__":
    args = sys.argv[1:]
    data = f"{date.today():%d-%m-%Y}"
    if "-d" in args:
        k = args.index("-d")
        data = args[k + 1]
        del args[k:k + 2]
    lista, fontes = args[0], [tuple(a.split("=", 1)) for a in args[1:]]
    for ol, n, arq in pedidos_ol(lista, data):
        print(f"OL  {ol}: {n} itens -> {arq}")
    if fontes:
        for nome, n, tot, arq in pedidos_distribuidor(*monta(lista, fontes), data):
            print(f"DIST {nome}: {n} itens, R$ {tot:,.2f} -> {arq}")
